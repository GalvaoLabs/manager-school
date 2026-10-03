"""Interface web simples para consultar e registrar notas escolares."""
import hmac
import ipaddress
import os
import secrets
from io import BytesIO
from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import pandas as pd
from flask import Flask, Response, abort, flash, redirect, render_template, request, send_file, session, url_for
from flask_limiter import Limiter
from openpyxl import load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Font, PatternFill
from matplotlib.backends.backend_pdf import PdfPages


def carregar_arquivo_env():
    """Lê pares simples CHAVE=VALOR do .env sem substituir variáveis existentes."""
    arquivo_env = Path(__file__).with_name(".env")
    if not arquivo_env.exists():
        return
    for linha in arquivo_env.read_text(encoding="utf-8").splitlines():
        linha = linha.strip()
        if not linha or linha.startswith("#") or "=" not in linha:
            continue
        chave, valor = linha.split("=", 1)
        os.environ.setdefault(chave.strip(), valor.strip().strip("\"'"))


carregar_arquivo_env()

from database import (
    cadastrar_aluno,
    consultar_dados_relatorio,
    criar_banco,
    excluir_aluno,
    excluir_nota,
    listar_alunos,
    listar_turmas,
    obter_aluno,
    obter_ou_criar_avaliacao,
    obter_ou_criar_disciplina,
    obter_ou_criar_oferta,
    obter_ou_criar_professor,
    obter_ou_criar_turma,
    consultar_notas_do_aluno,
    consultar_medias_por_materia_e_bimestre,
    salvar_nota,
    salvar_dados_do_aluno,
    atualizar_nota,
)

MODO_DEMO = os.environ.get("DEMO_MODE", "0") == "1"

app = Flask(__name__)
app.secret_key = os.environ.get("SECRET_KEY") or secrets.token_hex(32)
app.config.update(
    SESSION_COOKIE_HTTPONLY=True,
    SESSION_COOKIE_SAMESITE="Lax",
    SESSION_COOKIE_SECURE=os.environ.get("COOKIE_SECURE", "0") == "1",
)
USUARIO_APP = os.environ.get("APP_USERNAME", "admin")
SENHA_APP = "" if MODO_DEMO else os.environ.get("APP_PASSWORD", "")


def identificar_cliente():
    """Usa o IP encaminhado pelo proxy confiável do Render quando habilitado."""
    if os.environ.get("TRUST_X_FORWARDED_FOR") == "1":
        ip_encaminhado = request.headers.get("X-Forwarded-For", "").split(",", 1)[0].strip()
        try:
            return str(ipaddress.ip_address(ip_encaminhado))
        except ValueError:
            pass
    return request.remote_addr or "desconhecido"


limiter = Limiter(
    key_func=identificar_cliente,
    app=app,
    default_limits=["300 per hour", "2000 per day"],
    storage_uri=os.environ.get("RATELIMIT_STORAGE_URI", "memory://"),
)


@app.context_processor
def fornecer_dados_de_seguranca():
    """Disponibiliza o token dos formulários e o estado do acesso às páginas."""
    token = session.get("_csrf_token")
    if token is None:
        token = secrets.token_urlsafe(32)
        session["_csrf_token"] = token
    return {
        "csrf_token": token,
        "login_ativo": bool(SENHA_APP),
        "usuario_autenticado": bool(session.get("autenticado")),
        "modo_demo": MODO_DEMO,
    }


@app.before_request
def proteger_paginas_e_formularios():
    """Exige login quando há senha configurada e valida envios de formulários."""
    if MODO_DEMO and request.method == "POST":
        abort(403, description="A demonstração é somente para consulta.")

    if request.endpoint not in {"login", "static", "healthz"} and SENHA_APP:
        if not session.get("autenticado"):
            return redirect(url_for("login", proximo=request.full_path))

    if request.method == "POST":
        token_sessao = session.get("_csrf_token", "")
        token_formulario = request.form.get("csrf_token", "")
        if not token_sessao or not hmac.compare_digest(token_sessao, token_formulario):
            abort(400, description="O formulário expirou. Atualize a página e tente novamente.")


@app.route("/login", methods=["GET", "POST"])
@limiter.limit("5 per minute; 20 per hour", methods=["POST"], override_defaults=False)
def login():
    """Autentica o acesso público quando APP_PASSWORD está configurada."""
    if not SENHA_APP:
        return redirect(url_for("inicio"))
    if session.get("autenticado"):
        return redirect(url_for("inicio"))

    if request.method == "POST":
        usuario = request.form.get("usuario", "")
        senha = request.form.get("senha", "")
        usuario_valido = hmac.compare_digest(usuario, USUARIO_APP)
        senha_valida = hmac.compare_digest(senha, SENHA_APP)
        if usuario_valido and senha_valida:
            session.clear()
            session["autenticado"] = True
            destino = request.form.get("proximo", "")
            if not destino.startswith("/") or destino.startswith("//"):
                destino = url_for("inicio")
            return redirect(destino)
        flash("Usuário ou senha incorretos.", "erro")

    return render_template("login.html", proximo=request.args.get("proximo", ""))


@app.errorhandler(429)
def limite_de_acesso_excedido(_erro):
    """Explica em português quando o limite temporário de acessos é atingido."""
    return Response(
        "Muitas solicitações em pouco tempo. Aguarde alguns minutos e tente novamente.",
        status=429,
        mimetype="text/plain",
    )


@app.errorhandler(403)
def demonstracao_somente_consulta(_erro):
    if MODO_DEMO:
        return Response(
            "Esta demonstração é somente para consulta; os dados não podem ser alterados.",
            status=403,
            mimetype="text/plain",
        )
    return Response("Acesso não permitido.", status=403, mimetype="text/plain")


@app.post("/logout")
def logout():
    """Encerra a sessão autenticada."""
    session.clear()
    return redirect(url_for("login"))


@app.get("/healthz")
def healthz():
    """Endpoint simples usado pela hospedagem para verificar a aplicação."""
    return {"status": "ok"}


def montar_relatorio():
    """Carrega as notas em uma tabela pandas para facilitar a análise."""
    colunas = [
        "aluno_id", "aluno", "turma_id", "turma", "ano",
        "disciplina", "bimestre", "nota",
    ]
    dados = pd.DataFrame(consultar_dados_relatorio(), columns=colunas)
    return dados


def normalizar_id(valor):
    """Confere IDs vindos da URL e ignora valores inválidos."""
    try:
        numero = int(valor)
        return str(numero) if numero > 0 else ""
    except (TypeError, ValueError):
        return ""


def filtrar_relatorio(turma_id="", aluno_id=""):
    """Aplica os filtros escolhidos na página de relatórios."""
    dados = montar_relatorio()
    turma_id = normalizar_id(turma_id)
    aluno_id = normalizar_id(aluno_id)
    if turma_id:
        dados = dados[dados["turma_id"] == int(turma_id)]
    if aluno_id:
        dados = dados[dados["aluno_id"] == int(aluno_id)]
    return dados


def resumo_por_disciplina(dados):
    """Calcula a média e a quantidade de notas de cada período."""
    if dados.empty:
        return pd.DataFrame(columns=["disciplina", "bimestre", "media", "quantidade"])
    resumo = dados.groupby(["disciplina", "bimestre"], as_index=False).agg(
        media=("nota", "mean"), quantidade=("nota", "count")
    )
    resumo["media"] = resumo["media"].round(2)
    return resumo


def nome_do_relatorio(turma_id, aluno_id):
    """Monta um título de acordo com os filtros selecionados."""
    turma_id = normalizar_id(turma_id)
    aluno_id = normalizar_id(aluno_id)
    titulo = "Todas as turmas"
    if turma_id:
        turma = next((item for item in listar_turmas() if item[0] == int(turma_id)), None)
        if turma:
            titulo = f"Turma {turma[1]} ({turma[2]})"
    if aluno_id:
        aluno = obter_aluno(int(aluno_id))
        if aluno:
            titulo = f"{aluno[1]} — {aluno[2]} ({aluno[3]})"
    return titulo


def gerar_excel(dados, titulo):
    """Cria um Excel com dados detalhados, médias e gráfico de colunas."""
    arquivo = BytesIO()
    nomes_colunas = {
        "aluno_id": "ID do aluno", "aluno": "Aluno", "turma": "Turma",
        "ano": "Ano letivo", "disciplina": "Disciplina",
        "bimestre": "Bimestre", "nota": "Nota",
    }
    notas = dados[list(nomes_colunas)].rename(columns=nomes_colunas)
    medias = dados.groupby(
        ["aluno_id", "aluno", "disciplina", "bimestre"], as_index=False
    ).agg(media=("nota", "mean")) if not dados.empty else pd.DataFrame(
        columns=["aluno_id", "aluno", "disciplina", "bimestre", "media"]
    )
    medias["media"] = medias["media"].round(2)
    resumo = resumo_por_disciplina(dados)
    grafico_dados = resumo.pivot(
        index="disciplina", columns="bimestre", values="media"
    ) if not resumo.empty else pd.DataFrame()
    if not grafico_dados.empty:
        grafico_dados.columns = [f"{int(coluna)}º bimestre" for coluna in grafico_dados.columns]
        grafico_dados = grafico_dados.reset_index()
    else:
        grafico_dados = pd.DataFrame(columns=["Disciplina"])

    with pd.ExcelWriter(arquivo, engine="openpyxl") as escritor:
        notas.to_excel(escritor, sheet_name="Notas", index=False)
        medias.to_excel(escritor, sheet_name="Médias por aluno", index=False)
        grafico_dados.to_excel(escritor, sheet_name="Gráfico", index=False)

    arquivo.seek(0)
    livro = load_workbook(arquivo)
    folha = livro["Gráfico"]
    if folha.max_row > 1 and folha.max_column > 1:
        grafico = BarChart()
        grafico.type = "col"
        grafico.style = 10
        grafico.title = f"Médias — {titulo}"
        grafico.y_axis.title = "Média"
        grafico.x_axis.title = "Disciplina"
        grafico.y_axis.scaling.min = 0
        grafico.y_axis.scaling.max = 10
        valores = Reference(
            folha, min_col=2, max_col=folha.max_column,
            min_row=1, max_row=folha.max_row,
        )
        disciplinas = Reference(folha, min_col=1, min_row=2, max_row=folha.max_row)
        grafico.add_data(valores, titles_from_data=True)
        grafico.set_categories(disciplinas)
        grafico.height = 9
        grafico.width = 19
        folha.add_chart(grafico, "G2")

    for planilha in livro.worksheets:
        planilha.freeze_panes = "A2"
        planilha.auto_filter.ref = planilha.dimensions
        for celula in planilha[1]:
            celula.font = Font(bold=True, color="FFFFFF")
            celula.fill = PatternFill("solid", fgColor="6558E8")
        for coluna in planilha.columns:
            letra = coluna[0].column_letter
            maior = max((len(str(celula.value or "")) for celula in coluna), default=10)
            planilha.column_dimensions[letra].width = min(maior + 3, 38)

    arquivo_final = BytesIO()
    livro.save(arquivo_final)
    arquivo_final.seek(0)
    return arquivo_final


def gerar_pdf(dados, titulo):
    """Gera um PDF com gráfico, indicadores e uma tabela de médias."""
    arquivo = BytesIO()
    resumo = resumo_por_disciplina(dados)
    with PdfPages(arquivo) as pdf:
        figura, eixo = plt.subplots(figsize=(11.7, 8.3))
        figura.suptitle("Relatório de desempenho", fontsize=21, fontweight="bold")
        figura.text(0.5, 0.92, titulo, ha="center", color="#69758b", fontsize=12)
        alunos = dados["aluno_id"].nunique() if not dados.empty else 0
        media_geral = dados["nota"].mean() if not dados.empty else 0
        figura.text(
            0.5, 0.86,
            f"{alunos} aluno(s)   •   {len(dados)} nota(s)   •   média geral {media_geral:.2f}",
            ha="center", fontsize=12, color="#6558e8",
        )
        if dados.empty:
            eixo.text(0.5, 0.5, "Não há notas para este filtro.", ha="center", va="center")
            eixo.set_axis_off()
        else:
            tabela_grafico = dados.groupby(
                ["disciplina", "bimestre"]
            )["nota"].mean().unstack()
            tabela_grafico.plot(kind="bar", ax=eixo, color=["#6558e8", "#18a999", "#f2a541", "#e56b6f"])
            eixo.set_title("Média por disciplina e bimestre", pad=15)
            eixo.set_xlabel("Disciplina")
            eixo.set_ylabel("Média")
            eixo.set_ylim(0, 10)
            eixo.legend(title="Bimestre", ncol=4, loc="upper center", bbox_to_anchor=(0.5, -0.14))
            eixo.grid(axis="y", alpha=0.2)
            eixo.tick_params(axis="x", rotation=0)
        figura.tight_layout(rect=(0.05, 0.08, 0.95, 0.82))
        pdf.savefig(figura, bbox_inches="tight")
        plt.close(figura)

        if not resumo.empty:
            colunas = ["Disciplina", "Bimestre", "Média", "Notas consideradas"]
            linhas = resumo[["disciplina", "bimestre", "media", "quantidade"]].values.tolist()
            for inicio in range(0, len(linhas), 24):
                parte = linhas[inicio:inicio + 24]
                figura, eixo = plt.subplots(figsize=(11.7, 8.3))
                eixo.set_title(f"Médias por período — {titulo}", loc="left", pad=18, fontsize=17)
                eixo.axis("off")
                tabela = eixo.table(
                    cellText=[[d, f"{int(b)}º", f"{m:.2f}", int(q)] for d, b, m, q in parte],
                    colLabels=colunas, loc="center", cellLoc="left",
                )
                tabela.auto_set_font_size(False)
                tabela.set_fontsize(10)
                tabela.scale(1, 1.7)
                for (linha, _), celula in tabela.get_celld().items():
                    if linha == 0:
                        celula.set_facecolor("#6558e8")
                        celula.set_text_props(color="white", weight="bold")
                pdf.savefig(figura, bbox_inches="tight")
                plt.close(figura)
    arquivo.seek(0)
    return arquivo


@app.get("/")
def inicio():
    alunos = listar_alunos()
    return render_template("index.html", alunos=alunos)


@app.post("/notas")
def registrar_nota():
    aluno_id = request.form.get("aluno_id", "novo")
    disciplina = request.form.get("disciplina", "").strip()
    professor = request.form.get("professor", "").strip()
    titulo = request.form.get("avaliacao", "").strip()

    try:
        if not disciplina or not professor or not titulo:
            raise ValueError("Preencha disciplina, professor e avaliação.")

        bimestre = int(request.form.get("bimestre", ""))
        nota = float(request.form.get("nota", ""))
        nota_maxima = float(request.form.get("nota_maxima", "10"))
        if bimestre not in (1, 2, 3, 4):
            raise ValueError("Escolha um bimestre de 1 a 4.")
        if nota_maxima <= 0:
            raise ValueError("A nota máxima precisa ser maior que zero.")
        if not 0 <= nota <= nota_maxima:
            raise ValueError(f"A nota deve ficar entre 0 e {nota_maxima}.")

        if aluno_id == "novo":
            nome = request.form.get("nome", "").strip()
            turma = request.form.get("turma", "").strip()
            ano = int(request.form.get("ano", ""))
            if not nome or not turma:
                raise ValueError("Preencha o nome e a turma do novo aluno.")
            turma_id = obter_ou_criar_turma(turma, ano)
            aluno_id = cadastrar_aluno(nome, turma_id)
        else:
            aluno_id = int(aluno_id)
            aluno = obter_aluno(aluno_id)
            if aluno is None:
                raise ValueError("O aluno selecionado não existe.")
            turma_id = aluno[4]

        professor_id = obter_ou_criar_professor(professor)
        disciplina_id = obter_ou_criar_disciplina(disciplina)
        oferta_id = obter_ou_criar_oferta(turma_id, disciplina_id, professor_id)
        avaliacao_id = obter_ou_criar_avaliacao(
            oferta_id, titulo, bimestre, nota_maxima
        )
        salvar_nota(avaliacao_id, aluno_id, nota, nota_maxima)
        flash("Nota salva com sucesso.", "sucesso")
        return redirect(url_for("aluno", aluno_id=aluno_id))
    except (ValueError, TypeError) as erro:
        flash(str(erro), "erro")
        return redirect(url_for("inicio"))


@app.get("/alunos/<int:aluno_id>")
def aluno(aluno_id):
    dados_aluno = obter_aluno(aluno_id)
    if dados_aluno is None:
        flash("Aluno não encontrado.", "erro")
        return redirect(url_for("inicio"))
    notas = consultar_notas_do_aluno(aluno_id)
    medias = consultar_medias_por_materia_e_bimestre(aluno_id)
    return render_template("aluno.html", aluno=dados_aluno, notas=notas, medias=medias)


@app.post("/notas/<int:nota_id>/editar")
def editar_nota(nota_id):
    try:
        nova_nota = float(request.form.get("nota", ""))
        atualizar_nota(nota_id, nova_nota)
        flash("Nota atualizada.", "sucesso")
    except (ValueError, TypeError) as erro:
        flash(str(erro), "erro")
    aluno_id = int(request.form.get("aluno_id", "0"))
    return redirect(url_for("aluno", aluno_id=aluno_id))


@app.post("/notas/<int:nota_id>/excluir")
def remover_nota(nota_id):
    aluno_id = int(request.form.get("aluno_id", "0"))
    excluir_nota(nota_id)
    flash("Nota removida.", "sucesso")
    return redirect(url_for("aluno", aluno_id=aluno_id))


@app.post("/alunos/<int:aluno_id>/excluir")
def remover_aluno(aluno_id):
    excluir_aluno(aluno_id)
    flash("Aluno e notas relacionadas removidos.", "sucesso")
    return redirect(url_for("inicio"))


@app.get("/relatorios")
def relatorios():
    turma_id = normalizar_id(request.args.get("turma_id", ""))
    aluno_id = normalizar_id(request.args.get("aluno_id", ""))

    # O aluno é um filtro mais específico: ao selecioná-lo, usamos a turma
    # em que ele está cadastrado para manter os dois filtros coerentes.
    if aluno_id:
        aluno_selecionado = obter_aluno(int(aluno_id))
        if aluno_selecionado:
            turma_id = str(aluno_selecionado[4])

    dados = filtrar_relatorio(turma_id, aluno_id)
    alunos = listar_alunos(int(turma_id)) if turma_id else listar_alunos()
    resumo = resumo_por_disciplina(dados)
    tabela = resumo.to_dict("records")
    media_geral = round(float(dados["nota"].mean()), 2) if not dados.empty else 0
    return render_template(
        "relatorios.html",
        tabela=tabela,
        tem_dados=not dados.empty,
        turmas=listar_turmas(),
        alunos=alunos,
        turma_selecionada=turma_id,
        aluno_selecionado=aluno_id,
        titulo_relatorio=nome_do_relatorio(turma_id, aluno_id),
        quantidade_alunos=int(dados["aluno_id"].nunique()) if not dados.empty else 0,
        quantidade_notas=len(dados),
        media_geral=media_geral,
    )


@app.get("/relatorios/grafico.png")
def grafico_relatorio():
    dados = filtrar_relatorio(
        request.args.get("turma_id", ""), request.args.get("aluno_id", "")
    )
    tema_escuro = request.args.get("theme") == "dark"
    fundo_grafico = "#202622" if tema_escuro else "#F1EFE7"
    cor_texto = "#F0EFE8" if tema_escuro else "#292B26"
    cores_bimestres = (
        ["#E7A080", "#D78A68", "#C97554", "#F2C1A5"]
        if tema_escuro
        else ["#93432C", "#B36548", "#CC8A6E", "#E0B4A0"]
    )
    figura, eixo = plt.subplots(figsize=(9, 4.5))
    figura.patch.set_facecolor(fundo_grafico)
    eixo.set_facecolor(fundo_grafico)
    if dados.empty:
        eixo.text(0.5, 0.5, "Ainda não há notas", ha="center", va="center", color=cor_texto)
        eixo.set_axis_off()
    else:
        resumo = dados.groupby(["disciplina", "bimestre"])["nota"].mean().unstack(fill_value=0)
        resumo.plot(kind="bar", ax=eixo, color=cores_bimestres)
        eixo.set_title("Média por disciplina e bimestre", color=cor_texto)
        eixo.set_xlabel("Disciplina", color=cor_texto)
        eixo.set_ylabel("Média das notas", color=cor_texto)
        eixo.set_ylim(0, 10)
        eixo.tick_params(colors=cor_texto)
        eixo.grid(axis="y", color="#39423D" if tema_escuro else "#D8D6CB", alpha=0.65)
        for lado in eixo.spines.values():
            lado.set_color("#39423D" if tema_escuro else "#D8D6CB")
        eixo.legend(title="Bimestre", facecolor=fundo_grafico, labelcolor=cor_texto)
        figura.tight_layout()
    imagem = BytesIO()
    figura.savefig(imagem, format="png", bbox_inches="tight", facecolor=figura.get_facecolor())
    plt.close(figura)
    imagem.seek(0)
    return send_file(imagem, mimetype="image/png")


@app.get("/relatorios/csv")
def baixar_csv():
    dados = filtrar_relatorio(
        request.args.get("turma_id", ""), request.args.get("aluno_id", "")
    )
    conteudo = dados.to_csv(index=False).encode("utf-8-sig")
    return Response(
        conteudo,
        mimetype="text/csv; charset=utf-8",
        headers={"Content-Disposition": "attachment; filename=relatorio_notas.csv"},
    )


@app.get("/relatorios/excel")
def baixar_excel():
    turma_id = request.args.get("turma_id", "")
    aluno_id = request.args.get("aluno_id", "")
    dados = filtrar_relatorio(turma_id, aluno_id)
    arquivo = gerar_excel(dados, nome_do_relatorio(turma_id, aluno_id))
    return send_file(
        arquivo,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        as_attachment=True,
        download_name="relatorio_notas.xlsx",
    )


@app.get("/relatorios/pdf")
def baixar_pdf():
    turma_id = request.args.get("turma_id", "")
    aluno_id = request.args.get("aluno_id", "")
    dados = filtrar_relatorio(turma_id, aluno_id)
    arquivo = gerar_pdf(dados, nome_do_relatorio(turma_id, aluno_id))
    return send_file(
        arquivo,
        mimetype="application/pdf",
        as_attachment=True,
        download_name="relatorio_notas.pdf",
    )


criar_banco()
if MODO_DEMO:
    alunos_atuais = listar_alunos()
    if not alunos_atuais:
        from demo_data import popular_dados_demo

        popular_dados_demo(salvar_dados_do_aluno)


if __name__ == "__main__":
    app.run(debug=os.environ.get("FLASK_DEBUG", "1") == "1")

